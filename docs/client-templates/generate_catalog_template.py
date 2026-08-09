"""Build the client-facing catalog workbook.

    python docs/client-templates/generate_catalog_template.py

Writes `vista-store-catalog-template.xlsx` beside this file. The column names come from
the preparation pipeline (`backend/app/catalog_prep/prepare.py`) — this script only
arranges them for a non-technical owner: Arabic notes, dropdowns, frozen headers.

Sheet names must stay `categories` and `products`; every other sheet is ignored by the
parser, which is what keeps `Instructions` and the `*_example` sheets harmless.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).with_name("vista-store-catalog-template.xlsx")
DATA_ROWS = 300  # rows the dropdowns cover; blank rows are ignored by the parser

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF3CD")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F3A5F")
BODY_FONT = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", size=11, bold=True)
THIN = Side(style="thin", color="BFC9D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BOOL_HELP = "نعم أو لا (اختر من القائمة المنسدلة)"

# (column, width, Arabic note shown as a cell comment on the header)
CATEGORY_COLUMNS = [
    ("name", 26, "اسم التصنيف كما يظهر للزبون. مثال: هدايا الأعراس"),
    ("slug", 24, "معرّف إنجليزي بحروف صغيرة وشرطات، بدون مسافات. مثال: wedding-gifts\nمطلوب لأن الاسم عربي. يجب ألا يتكرر."),
    ("parent_name", 24, "اسم التصنيف الأب كما كُتب في عمود name بصف آخر من هذه الورقة.\nاتركه فارغًا للتصنيف الرئيسي."),
    ("description", 46, "وصف قصير للتصنيف (اختياري)."),
    ("image_filename", 26, "اسم ملف الصورة فقط، مثال: category-wedding.jpg\nليس رابطًا وليس مسارًا. تُرسل الصور في مجلد منفصل."),
    ("sort_order", 12, "رقم ترتيب الظهور (0، 1، 2 …). اتركه فارغًا ليُستخدم ترتيب الصفوف."),
    ("is_featured", 14, "هل يظهر التصنيف في الواجهة الرئيسية؟ " + BOOL_HELP),
    ("is_active", 12, "هل التصنيف مفعّل؟ " + BOOL_HELP),
]

PRODUCT_COLUMNS = [
    ("name", 30, "اسم المنتج بالعربية كما يظهر للزبون."),
    ("category_name", 24, "اسم التصنيف تمامًا كما كُتب في عمود name بورقة categories."),
    ("sku", 18, "رمز المنتج، فريد لكل منتج. مثال: VST-1001\nمطلوب: تُبنى عليه أسماء الصور والباقات."),
    ("product_type", 16, "standard = منتج عادي\npackage = باقة مكوّنة من منتجات أخرى\nsilicone_mold = قالب سيليكون"),
    ("price", 12, "السعر بالأرقام فقط، أكبر من صفر. مثال: 45 أو 45.50\nبدون رمز عملة وبدون فواصل."),
    ("compare_at_price", 16, "السعر قبل الخصم. يجب أن يكون أعلى من price، وإلا اتركه فارغًا."),
    ("cost_price", 12, "سعر التكلفة (داخلي، لا يظهر للزبون). اختياري."),
    ("stock_quantity", 14, "الكمية المتوفرة بالمخزون. رقم صحيح."),
    ("low_stock_threshold", 18, "الحد الذي يُعتبر عنده المخزون منخفضًا. اتركه فارغًا لاستخدام 3."),
    ("track_inventory", 16, "هل نتابع كمية المخزون لهذا المنتج؟ " + BOOL_HELP),
    ("short_description", 40, "سطر أو سطران يظهران في بطاقة المنتج."),
    ("description", 60, "الوصف الكامل للمنتج."),
    ("is_active", 12, "هل المنتج معروض في المتجر؟ " + BOOL_HELP),
    ("is_featured", 14, "منتج مميّز في الصفحة الرئيسية؟ " + BOOL_HELP),
    ("is_new", 12, "منتج جديد؟ " + BOOL_HELP),
    ("is_bestseller", 16, "الأكثر مبيعًا؟ " + BOOL_HELP),
    ("sort_order", 12, "رقم ترتيب الظهور داخل التصنيف. اتركه فارغًا ليُستخدم ترتيب الصفوف."),
    ("seo_title", 30, "عنوان صفحة المنتج في محركات البحث. اتركه فارغًا ليُشتق من الاسم."),
    ("seo_description", 46, "وصف الصفحة في محركات البحث. اختياري."),
    ("image_1", 22, "اسم ملف الصورة الرئيسية (صورة الغلاف). مثال: VST-1001-01.jpg"),
    ("image_2", 22, "الصورة الثانية. مثال: VST-1001-02.jpg\nوجودها يحسّن عرض المنتج في المتجر."),
    ("image_3", 22, "صورة إضافية. مثال: VST-1001-03.jpg"),
    ("image_4", 22, "صورة إضافية."),
    ("image_5", 22, "صورة إضافية. عند الحاجة لأكثر من ذلك أضف أعمدة image_6 وimage_7 بنفس التسمية."),
    ("spec_1_name", 18, "اسم المواصفة. مثال: المقاس"),
    ("spec_1_value", 22, "قيمة المواصفة. مثال: 20×20 سم"),
    ("spec_2_name", 18, "اسم المواصفة الثانية. مثال: الخامة"),
    ("spec_2_value", 22, "قيمة المواصفة الثانية. مثال: خشب طبيعي"),
    ("spec_3_name", 18, "اسم المواصفة الثالثة."),
    ("spec_3_value", 22, "قيمة المواصفة الثالثة."),
    ("spec_4_name", 18, "اسم المواصفة الرابعة."),
    ("spec_4_value", 22, "قيمة المواصفة الرابعة."),
    ("option_1_name", 18, "اسم خيار يختاره الزبون. مثال: اللون"),
    ("option_1_values", 26, "القيم مفصولة بفاصلة. مثال: أبيض، ذهبي، أسود"),
    ("option_2_name", 18, "اسم الخيار الثاني. مثال: المقاس"),
    ("option_2_values", 26, "القيم مفصولة بفاصلة. مثال: صغير، وسط، كبير"),
    ("package_items", 34, "للباقات فقط (product_type = package).\nرموز المنتجات والكميات: VST-1001 x2; VST-1005 x1\nكل رمز يجب أن يكون موجودًا في عمود sku بهذه الورقة."),
]

CATEGORY_EXAMPLES = [
    ["تصنيف تجريبي — هدايا", "example-gifts", "", "صف تجريبي: احذفه ولا تنسخه كما هو.", "EXAMPLE-CATEGORY-01.jpg", 1, "نعم", "نعم"],
    ["تصنيف تجريبي — قوالب", "example-molds", "تصنيف تجريبي — هدايا", "تصنيف فرعي تجريبي.", "EXAMPLE-CATEGORY-02.jpg", 2, "لا", "نعم"],
]

PRODUCT_EXAMPLES = [
    {
        "name": "منتج تجريبي — علبة هدايا",
        "category_name": "تصنيف تجريبي — هدايا",
        "sku": "EXAMPLE-001",
        "product_type": "standard",
        "price": 45,
        "compare_at_price": 60,
        "stock_quantity": 12,
        "track_inventory": "نعم",
        "short_description": "مثال لمنتج عادي بصورتين.",
        "description": "صف تجريبي للتوضيح فقط. احذف أوراق الأمثلة قبل الإرسال أو تجاهلها.",
        "is_active": "نعم",
        "is_featured": "نعم",
        "is_new": "نعم",
        "is_bestseller": "لا",
        "image_1": "EXAMPLE-001-01.jpg",
        "image_2": "EXAMPLE-001-02.jpg",
    },
    {
        "name": "منتج تجريبي — قالب سيليكون",
        "category_name": "تصنيف تجريبي — قوالب",
        "sku": "EXAMPLE-002",
        "product_type": "silicone_mold",
        "price": 30.5,
        "cost_price": 18,
        "stock_quantity": 40,
        "low_stock_threshold": 5,
        "track_inventory": "نعم",
        "short_description": "مثال لمنتج بمواصفات وخيارات.",
        "is_active": "نعم",
        "is_new": "لا",
        "image_1": "EXAMPLE-002-01.jpg",
        "spec_1_name": "المقاس",
        "spec_1_value": "20×20 سم",
        "spec_2_name": "الخامة",
        "spec_2_value": "سيليكون غذائي",
        "option_1_name": "اللون",
        "option_1_values": "أبيض، ذهبي، أسود",
        "option_2_name": "المقاس",
        "option_2_values": "صغير، كبير",
    },
    {
        "name": "منتج تجريبي — باقة",
        "category_name": "تصنيف تجريبي — هدايا",
        "sku": "EXAMPLE-003",
        "product_type": "package",
        "price": 99,
        "stock_quantity": 5,
        "track_inventory": "لا",
        "short_description": "مثال لباقة تحتوي منتجين.",
        "is_active": "نعم",
        "is_featured": "لا",
        "image_1": "EXAMPLE-003-01.jpg",
        "package_items": "EXAMPLE-001 x2; EXAMPLE-002 x1",
    },
]

INSTRUCTIONS = [
    ("title", "نموذج بيانات متجر Vista — التصنيفات والمنتجات"),
    ("text", "املأ ورقتَي categories و products فقط. باقي الأوراق للشرح والأمثلة ولا تُقرأ عند الاستيراد."),
    ("head", "قواعد عامة"),
    ("item", "لا تغيّر أسماء الأوراق: categories و products."),
    ("item", "لا تغيّر أسماء الأعمدة في الصف الأول، ولا تحذف أعمدة ولا تعيد ترتيبها بتغيير أسمائها."),
    ("item", "كل صف = تصنيف واحد أو منتج واحد."),
    ("item", "الخانة غير المطلوبة تُترك فارغة تمامًا. لا تكتب «-» ولا «لا يوجد» ولا «N/A»."),
    ("item", "مرّر الفأرة فوق عنوان أي عمود لترى شرحه (مثلث أحمر صغير في زاوية الخانة)."),
    ("head", "الأسعار والأرقام"),
    ("item", "الأسعار أرقام فقط: 45 أو 45.50 — بدون رمز عملة وبدون فواصل آلاف وبدون نص."),
    ("item", "price مطلوب لكل منتج ويجب أن يكون أكبر من صفر."),
    ("item", "compare_at_price هو السعر قبل الخصم، ويجب أن يكون أعلى من price، وإلا يُترك فارغًا."),
    ("head", "نعم / لا"),
    ("item", "كل خانات نعم/لا تُملأ من القائمة المنسدلة: نعم أو لا فقط."),
    ("item", "الأعمدة: is_active، is_featured، is_new، is_bestseller، track_inventory."),
    ("head", "رمز المنتج SKU"),
    ("item", "أدخل رمزًا فريدًا لكل منتج، مثل VST-1001. لا يتكرر رمز بين منتجين."),
    ("item", "الرمز مطلوب: تُبنى عليه أسماء ملفات الصور، وتُشير إليه الباقات."),
    ("head", "الصور"),
    ("item", "تُكتب الصور باسم الملف فقط، مثل VST-1001-01.jpg — بدون أي رابط أو مسار أو عنوان موقع."),
    ("item", "ترسل ملفات الصور نفسها في مجلد منفصل مع الملف، ولا تُلصق داخل الإكسل."),
    ("item", "التسمية الموصى بها: <SKU>-01.jpg و<SKU>-02.jpg و<SKU>-03.jpg."),
    ("item", "image_1 هي صورة الغلاف الرئيسية للمنتج، ثم image_2 وهكذا بالترتيب."),
    ("item", "وجود صورة ثانية يُحسّن عرض المنتج في المتجر بشكل ملحوظ."),
    ("item", "يفضَّل أن تكون صور المنتجات مربعة 1:1 وبمقاس 1000 بكسل على الأقل."),
    ("item", "يفضَّل أن تكون صور التصنيفات بنفس النسبة لجميع التصنيفات (مربعة 1:1) وبعرض 1000 بكسل على الأقل."),
    ("item", "أسماء الملفات حساسة لحالة الأحرف: VST-1001-01.jpg ليست vst-1001-01.jpg."),
    ("item", "إن احتجت أكثر من خمس صور، أضف أعمدة image_6 وimage_7 بنفس التسمية."),
    ("head", "التصنيفات"),
    ("item", "في ورقة categories: name هو الاسم العربي، وslug معرّف إنجليزي بحروف صغيرة وشرطات مثل wedding-gifts."),
    ("item", "التصنيف الفرعي يكتب اسم أبيه في parent_name تمامًا كما كُتب في عمود name."),
    ("item", "في ورقة products، عمود category_name يجب أن يطابق حرفيًا اسم التصنيف من ورقة categories."),
    ("head", "أنواع المنتجات (product_type)"),
    ("item", "standard: منتج عادي يُباع بمفرده."),
    ("item", "package: باقة مكوّنة من منتجات أخرى موجودة في نفس الورقة."),
    ("item", "silicone_mold: قالب سيليكون."),
    ("head", "الباقات (package_items)"),
    ("item", "تُملأ للباقات فقط، بصيغة: الرمز ثم x ثم الكمية، وتُفصل العناصر بفاصلة منقوطة."),
    ("item", "مثال: VST-1001 x2; VST-1005 x1"),
    ("item", "كل رمز مذكور يجب أن يكون موجودًا في عمود sku بورقة products."),
    ("head", "المواصفات والخيارات"),
    ("item", "المواصفات أزواج: spec_1_name مع spec_1_value، مثل: المقاس / 20×20 سم."),
    ("item", "الخيارات: option_1_name مثل «اللون»، وoption_1_values قيم مفصولة بفاصلة: أبيض، ذهبي، أسود."),
    ("head", "الأمثلة"),
    ("item", "أوراق categories_example وproducts_example للتوضيح فقط ولا تُستورد."),
    ("item", "لا تكتب بياناتك داخلها، ولا تنسخ صفوفها كما هي."),
]


def _style_header(sheet, columns) -> None:
    for index, (title, width, note) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        comment = Comment(note, "Vista Store")
        comment.width = 320
        comment.height = 110
        cell.comment = comment
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
    sheet.sheet_view.rightToLeft = True


def _validations(sheet, columns, last_row: int) -> None:
    boolean_columns = {"is_active", "is_featured", "is_new", "is_bestseller", "track_inventory"}
    names = [name for name, _, _ in columns]

    yes_no = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True)
    yes_no.error = "اكتب نعم أو لا فقط."
    yes_no.errorTitle = "قيمة غير مقبولة"
    sheet.add_data_validation(yes_no)

    for name in boolean_columns:
        if name not in names:
            continue
        letter = get_column_letter(names.index(name) + 1)
        yes_no.add(f"{letter}2:{letter}{last_row}")

    if "product_type" in names:
        types = DataValidation(
            type="list", formula1='"standard,package,silicone_mold"', allow_blank=True
        )
        types.error = "اختر standard أو package أو silicone_mold."
        types.errorTitle = "نوع غير معروف"
        sheet.add_data_validation(types)
        letter = get_column_letter(names.index("product_type") + 1)
        types.add(f"{letter}2:{letter}{last_row}")


def _wrap_text_columns(sheet, columns, last_row: int) -> None:
    wrapped = {"description", "short_description", "seo_description", "package_items"}
    names = [name for name, _, _ in columns]
    for name in wrapped:
        if name not in names:
            continue
        letter = get_column_letter(names.index(name) + 1)
        for row in range(2, last_row + 1):
            sheet[f"{letter}{row}"].alignment = Alignment(wrap_text=True, vertical="top")


def build_instructions(sheet) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 4
    sheet.column_dimensions["B"].width = 110

    row = 1
    for kind, text in INSTRUCTIONS:
        cell = sheet.cell(row=row, column=2, value=text)
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        if kind == "title":
            cell.font = TITLE_FONT
            sheet.row_dimensions[row].height = 28
            row += 2
            continue
        if kind == "head":
            cell.font = BOLD
            cell.fill = PatternFill("solid", fgColor="EAF0F6")
            row += 1
            continue
        cell.font = BODY_FONT
        if kind == "item":
            cell.value = f"• {text}"
        row += 1


def build_sheet(sheet, columns, *, rows: list[list] | None = None, example: bool = False) -> None:
    _style_header(sheet, columns)
    last_row = 1 + (len(rows) if rows else DATA_ROWS)
    for offset, values in enumerate(rows or [], start=2):
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=index, value=value)
            cell.border = BORDER
            if example:
                cell.fill = EXAMPLE_FILL
    _wrap_text_columns(sheet, columns, last_row)
    _validations(sheet, columns, last_row)


def product_rows(records: list[dict]) -> list[list]:
    names = [name for name, _, _ in PRODUCT_COLUMNS]
    return [[record.get(name) for name in names] for record in records]


def main() -> None:
    workbook = Workbook()
    build_instructions(workbook.active)
    workbook.active.title = "Instructions"

    build_sheet(workbook.create_sheet("categories"), CATEGORY_COLUMNS)
    build_sheet(workbook.create_sheet("products"), PRODUCT_COLUMNS)
    build_sheet(
        workbook.create_sheet("categories_example"), CATEGORY_COLUMNS,
        rows=CATEGORY_EXAMPLES, example=True,
    )
    build_sheet(
        workbook.create_sheet("products_example"), PRODUCT_COLUMNS,
        rows=product_rows(PRODUCT_EXAMPLES), example=True,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
