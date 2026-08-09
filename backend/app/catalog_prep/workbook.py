"""Reading the client's workbook into plain rows.

Two input shapes, because clients send both:

* an `.xlsx` workbook whose sheets are named `products` and `categories`;
* a directory holding `products.csv` and `categories.csv`.

Nothing here interprets a value. A cell arrives as whatever the file held — a string, a
number, a date, a bool, or None — and `prepare` decides what it means. Keeping the two
apart is what lets the error messages name a sheet, a row and a column.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

PRODUCTS_SHEET = "products"
CATEGORIES_SHEET = "categories"

XLSX_SUFFIXES = {".xlsx", ".xlsm"}

_HEADER_CLEAN = re.compile(r"[\s\-]+")


class WorkbookError(ValueError):
    """The input file cannot be read as a catalog workbook."""


@dataclass(frozen=True)
class SheetRow:
    """One data row, with its spreadsheet row number kept for error messages."""

    sheet: str
    number: int
    values: dict[str, Any]

    def raw(self, column: str) -> Any:
        return self.values.get(column)

    @property
    def is_empty(self) -> bool:
        return all(
            value is None or (isinstance(value, str) and not value.strip())
            for value in self.values.values()
        )


def normalise_header(value: Any) -> str:
    """`  Compare At Price ` -> `compare_at_price`."""
    text = str(value or "").strip().lower()
    text = _HEADER_CLEAN.sub("_", text)
    return re.sub(r"_{2,}", "_", text).strip("_")


def _rows_from_table(sheet: str, table: list[list[Any]]) -> list[SheetRow]:
    if not table:
        return []
    headers = [normalise_header(cell) for cell in table[0]]
    duplicates = sorted({h for h in headers if h and headers.count(h) > 1})
    if duplicates:
        raise WorkbookError(f"sheet {sheet!r} has duplicate column(s): {', '.join(duplicates)}")

    rows: list[SheetRow] = []
    for offset, raw_row in enumerate(table[1:], start=2):
        values = {
            header: (raw_row[position] if position < len(raw_row) else None)
            for position, header in enumerate(headers)
            if header
        }
        row = SheetRow(sheet=sheet, number=offset, values=values)
        if row.is_empty:
            continue  # trailing blank rows are what a spreadsheet always has
        rows.append(row)
    return rows


def _read_xlsx(path: Path) -> dict[str, list[SheetRow]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the install extra
        raise WorkbookError(
            "reading .xlsx needs openpyxl: pip install -e .[catalog-prep]"
        ) from exc

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheets: dict[str, list[SheetRow]] = {}
        for worksheet in workbook.worksheets:
            name = normalise_header(worksheet.title)
            if name not in (PRODUCTS_SHEET, CATEGORIES_SHEET):
                continue
            sheets[name] = _rows_from_table(name, [list(row) for row in worksheet.iter_rows(values_only=True)])
        return sheets
    finally:
        workbook.close()


def _read_csv_dir(path: Path) -> dict[str, list[SheetRow]]:
    sheets: dict[str, list[SheetRow]] = {}
    for name in (PRODUCTS_SHEET, CATEGORIES_SHEET):
        csv_path = path / f"{name}.csv"
        if not csv_path.is_file():
            continue
        # utf-8-sig: Excel writes a BOM, and a BOM in the first header would turn
        # `name` into something no column lookup would ever match.
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            sheets[name] = _rows_from_table(name, [list(row) for row in csv.reader(handle)])
    return sheets


def read_sheets(path: str | Path) -> dict[str, list[SheetRow]]:
    """Read `products` and `categories` from a workbook file or a directory of CSVs."""
    source = Path(path)
    if source.is_dir():
        sheets = _read_csv_dir(source)
    elif not source.is_file():
        raise WorkbookError(f"input not found: {source}")
    elif source.suffix.lower() in XLSX_SUFFIXES:
        sheets = _read_xlsx(source)
    elif source.suffix.lower() == ".csv":
        raise WorkbookError(
            f"{source.name} is a single CSV. Put products.csv and categories.csv in a "
            "directory and pass the directory, or supply one .xlsx workbook."
        )
    else:
        raise WorkbookError(f"unsupported input {source.name!r}: expected .xlsx or a directory")

    if PRODUCTS_SHEET not in sheets:
        raise WorkbookError(f"no {PRODUCTS_SHEET!r} sheet found in {source}")
    sheets.setdefault(CATEGORIES_SHEET, [])
    return sheets
